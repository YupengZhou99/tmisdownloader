#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
网页版 new 调整期国库收支存明细表批量下载工具
=============================================
菜单路径：固定报表 → 国库收支存明细表
iframe:   fineReportTsasRpt1010
导出流程：点击「导出」→「Excel」→「原样导出」

按附件 2 要求自动生成任务：
- 调整期年报 2003-2025
- 国库代码默认 1009000000
- 文件名 newszYYYY
- 导出后删除「年累计=0」行，以及科目 T040401/T050401 行

依赖：pip install playwright openpyxl psutil
      playwright install chromium
"""

import sys
import os
import re
import threading
import tkinter as tk
from tkinter import ttk, filedialog, scrolledtext, messagebox
from openpyxl import load_workbook
from datetime import datetime
from decimal import Decimal, InvalidOperation
import time
import traceback
import asyncio

from playwright.async_api import async_playwright, Page

try:
    import psutil
except ImportError:
    psutil = None

try:
    import pyperclip
except ImportError:
    pyperclip = None


# ============================================================================
# 常量
# ============================================================================
IFRAME_NAME  = "fineReportTsasRpt1010"
MENU_PATH    = ["固定报表", "国库收支存明细表"]
DEFAULT_TREASURY_CODE = "1009000000"
DEFAULT_START_YEAR = 2003
DEFAULT_END_YEAR = 2025
EXCLUDED_SUBJECT_CODES = {"T040401", "T050401"}

# 全局停止标志
stop_flag    = False
sniper_flag  = False


# ============================================================================
# 主应用类
# ============================================================================
class GDMXApp(tk.Tk):
    """网页版 new 调整期国库收支存明细表批量下载工具"""

    BG_PRIMARY   = "#1a1a2e"
    BG_CARD      = "#16213e"
    BG_INPUT     = "#0f3460"
    BG_LOG       = "#0d1117"
    FG_PRIMARY   = "#e8e8e8"
    FG_SECONDARY = "#8892a8"
    FG_ACCENT    = "#00d2ff"
    CLR_GREEN    = "#00e676"
    CLR_RED      = "#ff5252"
    CLR_YELLOW   = "#ffd740"
    CLR_BLUE     = "#448aff"
    CLR_HOVER_G  = "#00c853"
    CLR_HOVER_R  = "#ff1744"
    CLR_HOVER_B  = "#2979ff"

    def __init__(self):
        super().__init__()
        self.title("调整期国库收支存明细表  批量下载工具")
        self.geometry("980x820")
        self.resizable(True, True)
        self.configure(bg=self.BG_PRIMARY)
        self.minsize(800, 650)

        self.treasury_code_var = tk.StringVar(value=DEFAULT_TREASURY_CODE)
        self.download_folder   = tk.StringVar()
        self.start_year_var    = tk.StringVar(value=str(DEFAULT_START_YEAR))
        self.end_year_var      = tk.StringVar(value=str(DEFAULT_END_YEAR))
        self.enable_postprocess = tk.BooleanVar(value=True)
        self.chrome_path_var   = tk.StringVar()
        self.target_url        = None
        self.is_running        = False
        self.worker_thread     = None
        self.sniper_thread     = None
        self.current_nav_done  = False

        self._init_styles()
        self._build_gui()

    # ------------------------------------------------------------------ styles
    def _init_styles(self):
        style = ttk.Style(self)
        style.theme_use("clam")
        style.configure("Card.TFrame",  background=self.BG_CARD)
        style.configure("Main.TFrame",  background=self.BG_PRIMARY)
        style.configure("Title.TLabel", background=self.BG_PRIMARY,
                        foreground=self.FG_ACCENT,  font=("Segoe UI", 16, "bold"))
        style.configure("Sub.TLabel",   background=self.BG_PRIMARY,
                        foreground=self.FG_SECONDARY, font=("Segoe UI", 9))
        style.configure("Card.TLabel",  background=self.BG_CARD,
                        foreground=self.FG_SECONDARY, font=("Segoe UI", 10))
        style.configure("LogTitle.TLabel", background=self.BG_PRIMARY,
                        foreground=self.FG_SECONDARY, font=("Segoe UI", 9, "bold"))
        style.configure("Path.TEntry",
                        fieldbackground=self.BG_INPUT, foreground=self.FG_PRIMARY,
                        borderwidth=0, padding=(8, 6))
        style.map("Path.TEntry",
                  fieldbackground=[("readonly", self.BG_INPUT)],
                  foreground=[("readonly", self.FG_PRIMARY)])

    def _make_btn(self, parent, text, command, color, hover, width=16, state=tk.NORMAL):
        btn = tk.Button(
            parent, text=text, command=command,
            bg=color, fg="white", activebackground=hover, activeforeground="white",
            font=("Segoe UI", 10, "bold"), relief="flat", cursor="hand2",
            bd=0, padx=14, pady=8, width=width, state=state,
            highlightthickness=0
        )
        btn.bind("<Enter>", lambda e: btn.config(bg=hover)  if btn["state"] != "disabled" else None)
        btn.bind("<Leave>", lambda e: btn.config(bg=color)  if btn["state"] != "disabled" else None)
        btn._base_color  = color
        btn._hover_color = hover
        return btn

    # ------------------------------------------------------------------ GUI
    def _build_gui(self):
        c = ttk.Frame(self, style="Main.TFrame", padding=(24, 16))
        c.pack(fill=tk.BOTH, expand=True)

        # 标题
        hdr = ttk.Frame(c, style="Main.TFrame")
        hdr.pack(fill=tk.X, pady=(0, 12))
        ttk.Label(hdr, text="调整期国库收支存明细表  批量下载", style="Title.TLabel").pack(side=tk.LEFT)
        ttk.Label(hdr, text="固定报表 · 网页版 new · 年报 · 全辖",
                  style="Sub.TLabel").pack(side=tk.LEFT, padx=(14, 0), pady=(8, 0))

        # 文件卡片
        card = ttk.Frame(c, style="Card.TFrame", padding=(20, 16))
        card.pack(fill=tk.X, pady=(0, 10))

        self._input_row(card, "国库代码", self.treasury_code_var,
                        "用于单库核对，默认 1009000000", row_pady=(0, 10))
        self._file_row(card, "保存目录", self.download_folder, self._select_folder)

        # 年报行
        row_yr = ttk.Frame(card, style="Card.TFrame")
        row_yr.pack(fill=tk.X, pady=(10, 0))
        ttk.Label(row_yr, text="调整期年报", style="Card.TLabel", width=10).pack(side=tk.LEFT)
        ttk.Label(row_yr, text="起始年", style="Card.TLabel", width=6).pack(side=tk.LEFT, padx=(8, 0))
        tk.Entry(row_yr, textvariable=self.start_year_var,
                 bg=self.BG_INPUT, fg=self.FG_PRIMARY, insertbackground=self.FG_ACCENT,
                 font=("Segoe UI", 10), relief="flat", bd=0, highlightthickness=0, width=8
                 ).pack(side=tk.LEFT, padx=(6, 16))
        ttk.Label(row_yr, text="终止年", style="Card.TLabel", width=6).pack(side=tk.LEFT)
        tk.Entry(row_yr, textvariable=self.end_year_var,
                 bg=self.BG_INPUT, fg=self.FG_PRIMARY, insertbackground=self.FG_ACCENT,
                 font=("Segoe UI", 10), relief="flat", bd=0, highlightthickness=0, width=8
                 ).pack(side=tk.LEFT, padx=(6, 10))
        ttk.Label(row_yr, text="默认 2003-2025", style="Card.TLabel").pack(side=tk.LEFT, padx=(10, 0))

        opt_row = ttk.Frame(card, style="Card.TFrame")
        opt_row.pack(fill=tk.X, pady=(10, 0))
        tk.Checkbutton(
            opt_row,
            text="启用数据处理（删除年累计为0行 + 剔除 T040401/T050401）",
            variable=self.enable_postprocess,
            bg=self.BG_CARD, fg=self.FG_SECONDARY,
            activebackground=self.BG_CARD, activeforeground=self.FG_PRIMARY,
            selectcolor=self.BG_INPUT, font=("Segoe UI", 10),
            highlightthickness=0, bd=0
        ).pack(side=tk.LEFT)

        # Chrome 行
        row4 = ttk.Frame(card, style="Card.TFrame")
        row4.pack(fill=tk.X, pady=(10, 0))
        ttk.Label(row4, text="Chrome路径", style="Card.TLabel", width=10).pack(side=tk.LEFT)
        tk.Entry(row4, textvariable=self.chrome_path_var,
                 bg=self.BG_INPUT, fg=self.FG_PRIMARY, insertbackground=self.FG_ACCENT,
                 font=("Segoe UI", 10), relief="flat", bd=0, highlightthickness=0
                 ).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(8, 10))
        ttk.Label(row4, text="（空=自动检测）", style="Card.TLabel").pack(side=tk.LEFT)

        # 按钮栏
        btn_bar = ttk.Frame(c, style="Main.TFrame")
        btn_bar.pack(fill=tk.X, pady=(4, 10))

        self.start_btn = self._make_btn(btn_bar, "  开始批量下载  ", self._start_task,
                                        self.CLR_GREEN, self.CLR_HOVER_G)
        self.start_btn.pack(side=tk.LEFT, padx=(0, 8))

        self.stop_btn = self._make_btn(btn_bar, "  停止  ", self._stop_task,
                                       self.CLR_RED, self.CLR_HOVER_R, width=8, state=tk.DISABLED)
        self.stop_btn.pack(side=tk.LEFT, padx=(0, 8))

        self.sniper_btn = self._make_btn(btn_bar, "  捕获登录链接  ", self._start_sniper,
                                         self.CLR_BLUE, self.CLR_HOVER_B, width=16)
        self.sniper_btn.pack(side=tk.RIGHT)

        # 日志
        ttk.Label(c, text="CONSOLE", style="LogTitle.TLabel").pack(anchor=tk.W, pady=(0, 4))
        lf = tk.Frame(c, bg=self.BG_LOG, bd=0, highlightthickness=1, highlightbackground="#2a2a4a")
        lf.pack(fill=tk.BOTH, expand=True)
        self.log_text = tk.Text(lf, bg=self.BG_LOG, fg="#8cc8a8",
                                font=("Cascadia Code", 10), wrap=tk.WORD,
                                insertbackground=self.FG_ACCENT, selectbackground="#264f78",
                                relief="flat", bd=0, padx=12, pady=10, highlightthickness=0)
        sb = ttk.Scrollbar(lf, orient=tk.VERTICAL, command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=sb.set)
        sb.pack(side=tk.RIGHT, fill=tk.Y)
        self.log_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.log_text.tag_config("INFO",    foreground="#8cc8a8")
        self.log_text.tag_config("SUCCESS", foreground=self.CLR_GREEN)
        self.log_text.tag_config("ERROR",   foreground=self.CLR_RED)
        self.log_text.tag_config("WARN",    foreground=self.CLR_YELLOW)
        self.log_text.tag_config("TS",      foreground="#5a6a7a")
        self.log_text.config(state=tk.DISABLED)

    def _file_row(self, parent, label, var, cmd, row_pady=(0, 0)):
        row = ttk.Frame(parent, style="Card.TFrame")
        row.pack(fill=tk.X, pady=row_pady)
        ttk.Label(row, text=label, style="Card.TLabel", width=10).pack(side=tk.LEFT)
        ttk.Entry(row, textvariable=var, style="Path.TEntry", state="readonly"
                  ).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(8, 10))
        self._make_btn(row, "选择", cmd, self.BG_INPUT, "#1a5276", width=8).pack(side=tk.LEFT)

    def _input_row(self, parent, label, var, hint="", row_pady=(0, 0)):
        row = ttk.Frame(parent, style="Card.TFrame")
        row.pack(fill=tk.X, pady=row_pady)
        ttk.Label(row, text=label, style="Card.TLabel", width=10).pack(side=tk.LEFT)
        tk.Entry(row, textvariable=var,
                 bg=self.BG_INPUT, fg=self.FG_PRIMARY, insertbackground=self.FG_ACCENT,
                 font=("Segoe UI", 10), relief="flat", bd=0, highlightthickness=0
                 ).pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(8, 10))
        if hint:
            ttk.Label(row, text=hint, style="Card.TLabel").pack(side=tk.LEFT)

    # ------------------------------------------------------------------ actions
    def _select_folder(self):
        p = filedialog.askdirectory(title="选择保存目录")
        if p:
            self.download_folder.set(p)
            self.log(f"已选择保存目录: {p}")

    def log(self, msg: str, level: str = "INFO"):
        def _append():
            self.log_text.config(state=tk.NORMAL)
            ts = datetime.now().strftime("%H:%M:%S")
            self.log_text.insert(tk.END, f"[{ts}]", "TS")
            self.log_text.insert(tk.END, f" {msg}\n", level)
            self.log_text.see(tk.END)
            self.log_text.config(state=tk.DISABLED)
        self.after(0, _append)

    # ------------------------------------------------------------------ sniper
    def _start_sniper(self):
        global sniper_flag
        if psutil is None:
            messagebox.showerror("缺少依赖", "请先安装 psutil: pip install psutil")
            return
        if not self.download_folder.get() or not os.path.isdir(self.download_folder.get()):
            messagebox.showwarning("提示", "请先选择保存目录")
            return
        sniper_flag = False
        self.target_url = None
        self.sniper_btn.config(state=tk.DISABLED, text="  正在监控...  ", bg="#5c6bc0")
        self.log("URL 狙击手已启动，请在客户端触发系统登录...", "WARN")
        self.sniper_thread = threading.Thread(target=self._sniper_worker, daemon=True)
        self.sniper_thread.start()

    def _sniper_worker(self):
        global sniper_flag
        browser_names = {"chrome.exe", "msedge.exe", "firefox.exe",
                         "chrome", "msedge", "firefox"}
        known_pids = set()
        for proc in psutil.process_iter(["pid", "name"]):
            try:
                if (proc.info["name"] or "").lower() in browser_names:
                    known_pids.add(proc.info["pid"])
            except (psutil.NoSuchProcess, psutil.AccessDenied):
                pass

        captured_url = None
        start_time = time.time()
        while not sniper_flag and (time.time() - start_time) < 180:
            try:
                for proc in psutil.process_iter(["pid", "name", "cmdline"]):
                    try:
                        pname = (proc.info["name"] or "").lower()
                        if pname not in browser_names or proc.info["pid"] in known_pids:
                            continue
                        for arg in (proc.info["cmdline"] or []):
                            if isinstance(arg, str) and arg.startswith("http"):
                                captured_url = arg
                                try:
                                    proc.kill()
                                except Exception:
                                    pass
                                break
                        if captured_url:
                            break
                        known_pids.add(proc.info["pid"])
                    except (psutil.NoSuchProcess, psutil.AccessDenied):
                        continue
                if captured_url:
                    break
            except Exception as e:
                self.log(f"狙击手异常: {e}", "ERROR")
            time.sleep(0.01)

        if captured_url:
            self.target_url = captured_url
            self.log("Token 链接捕获成功！", "SUCCESS")
            if pyperclip:
                try:
                    pyperclip.copy(captured_url)
                except Exception:
                    pass
            self.after(0, lambda: self.sniper_btn.config(
                state=tk.NORMAL, text="  捕获登录链接  ", bg=self.CLR_BLUE))
            self.after(500, self._start_task)
        else:
            self.log("狙击手超时（180秒），未捕获到链接", "WARN")
            self.after(0, lambda: self.sniper_btn.config(
                state=tk.NORMAL, text="  捕获登录链接  ", bg=self.CLR_BLUE))

    # ------------------------------------------------------------------ task control
    def _start_task(self):
        global stop_flag
        treasury_code = self.treasury_code_var.get().strip()
        if not treasury_code:
            messagebox.showerror("错误", "请输入国库代码")
            return
        if not self.download_folder.get() or not os.path.isdir(self.download_folder.get()):
            messagebox.showerror("错误", "请选择有效的保存目录")
            return
        if not self.target_url:
            messagebox.showwarning("提示",
                "尚未捕获到登录链接。\n请先点击「捕获登录链接」，然后在客户端触发系统登录。")
            return
        stop_flag = False
        self.current_nav_done = False
        self.start_btn.config(state=tk.DISABLED)
        self.stop_btn.config(state=tk.NORMAL)
        self.is_running = True
        self.log("=" * 50, "SUCCESS")
        self.log("批量下载任务开始", "SUCCESS")
        self.log(f"国库代码: {treasury_code}", "INFO")
        self.log("=" * 50, "SUCCESS")
        self.worker_thread = threading.Thread(
            target=self._worker,
            args=(treasury_code, self.download_folder.get()),
            daemon=True
        )
        self.worker_thread.start()

    def _stop_task(self):
        global stop_flag
        stop_flag = True
        self.log("用户请求停止...", "WARN")

    def _worker(self, treasury_code, dl_folder):
        try:
            asyncio.run(self._run_automation(treasury_code, dl_folder))
        except Exception as e:
            self.log(f"任务出错: {e}", "ERROR")
            self.log(traceback.format_exc(), "ERROR")
        finally:
            self.is_running = False
            self.after(0, lambda: self.start_btn.config(state=tk.NORMAL))
            self.after(0, lambda: self.stop_btn.config(state=tk.DISABLED))
            self.log("=" * 50, "SUCCESS")
            self.log("全部任务执行完毕", "SUCCESS")
            self.log("=" * 50, "SUCCESS")

    # ================================================================
    # 核心自动化
    # ================================================================
    async def _run_automation(self, treasury_code: str, dl_folder: str):
        global stop_flag

        # 1. 单库核对：国库代码直接来自界面输入，辖属固定为全辖。
        treasury_code = str(treasury_code).strip()
        if not treasury_code:
            self.log("国库代码为空", "ERROR")
            return
        self.log(f"使用国库代码: {treasury_code}", "SUCCESS")

        # 2. 生成年份列表。PPT要求网页版调整期年报为 2003-2025。
        try:
            sy = int(self.start_year_var.get().strip())
            ey = int(self.end_year_var.get().strip())
            if sy > ey:
                raise ValueError("起始年不能晚于终止年")
            years = [str(y) for y in range(sy, ey + 1)]
            self.log(f"调整期年报范围: {years[0]} ~ {years[-1]}，共 {len(years)} 年")
        except Exception as e:
            self.log(f"年份解析失败: {e}", "ERROR")
            return

        # 3. 构建任务列表：调整期年报
        tasks = []
        for year in years:
            tasks.append({
                "rpt_type": "5",          # 5 -- 年
                "date":     year,         # YYYY
                "tre_code": treasury_code,
                "gov_code": "0",          # 0 -- 全辖
            })

        total = len(tasks)
        self.log(f"共生成 {total} 个调整期年报任务", "SUCCESS")

        # 4. 启动浏览器
        self.log("启动 Chrome 浏览器...")
        async with async_playwright() as pw:
            try:
                launch_kwargs = {
                    "headless": False,
                    "channel": "chrome",
                    "args": [
                        "--ignore-certificate-errors",
                        "--ignore-ssl-errors",
                        "--disable-web-security",
                        "--no-first-run",
                        "--disable-popup-blocking",
                        "--start-maximized",
                    ]
                }
                chrome_path = self.chrome_path_var.get().strip()
                if chrome_path and os.path.exists(chrome_path):
                    launch_kwargs.pop("channel", None)
                    launch_kwargs["executable_path"] = chrome_path
                browser = await pw.chromium.launch(**launch_kwargs)
            except Exception as e:
                self.log(f"启动浏览器失败: {e}", "ERROR")
                return

            context = await browser.new_context(
                ignore_https_errors=True,
                viewport=None,
                accept_downloads=True,
            )
            page = await context.new_page()

            self.log("携 Token 登录内网系统...")
            try:
                await page.goto(self.target_url, wait_until="networkidle", timeout=60000)
                self.log("登录成功", "SUCCESS")
            except Exception as e:
                self.log(f"页面加载提示: {e}", "WARN")

            await asyncio.sleep(3)

            # 5. 循环执行任务
            success_count = fail_count = 0

            for idx, task in enumerate(tasks):
                if stop_flag:
                    self.log("已停止", "WARN")
                    break

                rpt_type  = task["rpt_type"]
                date      = task["date"]
                tre_code  = task["tre_code"]
                gov_code  = task["gov_code"]
                rpt_label = "调整期年报"
                fname = f"newsz{date}"

                self.log(f"\n{'─'*45}")
                self.log(f"▶ [{idx+1}/{total}] {fname}（{rpt_label}）")
                self.log(f"{'─'*45}")

                try:
                    # 首次或菜单未打开时导航
                    if not self.current_nav_done:
                        await self._navigate_to_report(page)
                        self.current_nav_done = True
                    else:
                        self.log("  复用已打开的报表页面", "INFO")

                    await asyncio.sleep(1.5)
                    await self._fill_form(page, date, tre_code, gov_code, rpt_type)
                    await self._click_query_and_wait(page)
                    saved_path = await self._export_and_save(page, dl_folder, fname)
                    if self.enable_postprocess.get():
                        self._postprocess_adjustment_excel(saved_path)
                    else:
                        self.log("  数据处理已关闭，跳过调整期清理", "INFO")

                    success_count += 1
                    self.log(f"  已保存: {os.path.basename(saved_path)}", "SUCCESS")

                except Exception as e:
                    fail_count += 1
                    self.log(f"  任务失败: {e}", "ERROR")
                    self.log(traceback.format_exc(), "ERROR")
                    # 失败后重置导航状态，下一条任务重新导航
                    self.current_nav_done = False

                await asyncio.sleep(1.5)

            self.log(f"\n汇总: 共{total}条，成功{success_count}，失败{fail_count}",
                     "SUCCESS" if fail_count == 0 else "WARN")

            try:
                await context.close()
                await browser.close()
            except Exception:
                pass

    # ================================================================
    # 导航
    # ================================================================
    async def _navigate_to_report(self, page: Page):
        """点击侧边栏菜单：固定报表 → 国库收支存明细表"""
        self.log(f"  导航到菜单: {' → '.join(MENU_PATH)}")

        for i, item_text in enumerate(MENU_PATH):
            target = page.get_by_text(item_text, exact=True).first
            # 如果是父级菜单，检查子项是否已展开
            if i < len(MENU_PATH) - 1:
                next_item = page.get_by_text(MENU_PATH[i + 1], exact=True).first
                try:
                    if await next_item.is_visible(timeout=1000):
                        self.log(f"    子菜单已展开，跳过: {item_text}")
                        continue
                except Exception:
                    pass

            try:
                await target.wait_for(state="visible", timeout=10000)
                await target.click()
                await asyncio.sleep(1)
            except Exception:
                self.log(f"    补偿点击: {item_text}", "WARN")
                try:
                    await target.click(force=True)
                    await asyncio.sleep(1)
                except Exception as e2:
                    self.log(f"    菜单点击失败: {item_text} - {e2}", "ERROR")
                    raise

        self.log("  导航完成", "SUCCESS")

    # ================================================================
    # 表单填充
    # ================================================================
    async def _fill_form(self, page: Page, date: str, tre_code: str,
                         gov_code: str, rpt_type: str = "3"):
        """
        填写表单字段。
        PPT 核对场景固定为调整期年报：date 为 YYYY。
        """
        self.log(f"  填充表单（调整期年报，日期={date}）...")

        # 报表库数据源选择 → 1 -- 报表库
        await self._select_dropdown(page, "pRptDbType", "1")

        # 调整期标志 → 1 -- 调整期（页面HTML中 0=正常期，1=调整期）
        await self._select_dropdown(page, "pTrimFlag", "1")

        # 金额单位 → 0 -- 元
        await self._select_dropdown(page, "pAmtUnit", "0")

        # 报表类型 → 5 -- 年
        await self._select_dropdown(page, "pRptType", "5")

        # 日期选项 → YYYY
        await self._fill_date(page, "pDate", date, "5")

        # 国库选择
        await self._fill_treasury(page, "pTreCode", tre_code)

        # 计划单列市/县/乡 → 0 -- 不含
        await self._select_dropdown(page, "pSigTreArea", "0")

        # 辖属区标志 → 0 -- 全辖
        await self._select_dropdown(page, "pGovernFlag", gov_code)

        # 科目大类选择 → 4 -- 目
        await self._select_dropdown(page, "pStatSbtLevel", "4")

        self.log("  表单填充完成", "SUCCESS")

    async def _select_dropdown(self, page: Page, field_id: str, value: str):
        """Element-UI 下拉菜单选择（容器定位法）"""
        self.log(f"    下拉 {field_id} = {value}")
        form_item = page.locator(f'.el-form-item:visible:has(label[for="{field_id}"])').first
        inp = form_item.locator('.el-select .el-input__inner').first
        await inp.wait_for(state="visible", timeout=10000)
        await inp.click()
        await asyncio.sleep(0.6)

        found = False
        # Element-UI 下拉弹出层在全局 body 下，需用 page.locator
        # 等待下拉列表出现
        try:
            await page.wait_for_selector(
                ".el-select-dropdown:not([style*='display: none']) .el-select-dropdown__item",
                timeout=3000
            )
        except Exception:
            pass

        items = page.locator(
            ".el-select-dropdown:not([style*='display: none']) .el-select-dropdown__item"
        )
        count = await items.count()
        for i in range(count):
            item = items.nth(i)
            text = (await item.text_content() or "").strip()
            # 匹配策略：完整相等、以"值 "开头（如"0 -- 元"）、值在文本中
            if (text == value
                    or text.startswith(f"{value} ")
                    or text.startswith(f"{value}--")
                    or f"-- {value}" in text
                    or value in text):
                await item.click()
                found = True
                break

        if not found:
            # 回退：尝试 :has-text 选择器
            try:
                candidate = page.locator(
                    f".el-select-dropdown:not([style*='display: none']) "
                    f".el-select-dropdown__item:has-text('{value}')"
                ).first
                if await candidate.is_visible(timeout=1000):
                    await candidate.click()
                    found = True
            except Exception:
                pass

        if not found:
            self.log(f"    下拉选项 '{value}' 未找到，跳过", "WARN")
            await page.keyboard.press("Escape")

        await asyncio.sleep(0.3)

    async def _fill_date(self, page: Page, field_id: str, date: str, rpt_type: str = "3"):
        """
        填充日期选择器。
        调整期年报：date=YYYY，直接输入如 2024。
        """
        self.log(f"    日期 {field_id} = {date}（年报）")
        display_val = date                            # 直接输入，无需转换格式

        form_item = page.locator(f'.el-form-item:visible:has(label[for="{field_id}"])').first
        inp = form_item.locator('.el-date-editor .el-input__inner').first
        await inp.wait_for(state="visible", timeout=10000)

        await inp.click()
        await asyncio.sleep(0.3)
        await inp.press("Control+a")
        await asyncio.sleep(0.1)
        await inp.press("Delete")
        await asyncio.sleep(0.1)
        await inp.type(display_val, delay=60)
        await asyncio.sleep(0.4)
        await inp.press("Enter")
        await asyncio.sleep(0.5)
        # 关闭可能残留的日历弹层
        try:
            picker = page.locator(".el-month-table, .el-year-table, .el-picker-panel")
            if await picker.first.is_visible(timeout=500):
                await page.keyboard.press("Escape")
                await asyncio.sleep(0.2)
        except Exception:
            pass

    async def _fill_treasury(self, page: Page, field_id: str, tre_code: str):
        """在国库选择输入框中直接键入国库代码，并触发 Vue 数据绑定"""
        self.log(f"    国库 {field_id} = {tre_code}")
        form_item = page.locator(f'.el-form-item:visible:has(label[for="{field_id}"])').first
        inp = form_item.locator('.el-input__inner').first
        await inp.wait_for(state="visible", timeout=10000)

        await inp.click()
        await asyncio.sleep(0.3)
        await inp.press("Control+a")       # 全选已有内容
        await asyncio.sleep(0.1)
        await inp.press("Delete")          # 清空
        await asyncio.sleep(0.1)
        await inp.type(tre_code, delay=40)
        await asyncio.sleep(0.3)
        # 触发 input/change 事件让 Vue 更新数据绑定
        await inp.dispatch_event("input")
        await asyncio.sleep(0.1)
        await inp.dispatch_event("change")
        await asyncio.sleep(0.1)
        await inp.press("Enter")
        await asyncio.sleep(0.4)
        # Tab 跳出字段，确保 blur 事件触发
        await inp.press("Tab")
        await asyncio.sleep(0.3)

    # ================================================================
    # 查询等待
    # ================================================================
    async def _click_query_and_wait(self, page: Page):
        """点击查询，等待 iframe 内报表渲染完成"""
        self.log("  点击查询...")
        query_btn = page.locator("button.el-button:visible:has-text('查询')").first
        await query_btn.wait_for(state="visible", timeout=5000)
        await query_btn.click()

        self.log("  等待报表加载（最长 6 分钟）...")
        await asyncio.sleep(3)

        iframe = page.frame_locator(f'iframe[name="{IFRAME_NAME}"]')
        # 以「导出」按钮（button.fr-btn-text.x-emb-export）的父级 .fr-btn 是否含
        # ui-state-enabled 作为报表加载完成的信号
        max_wait, poll, elapsed = 360, 2, 0
        export_ready = False
        while elapsed < max_wait:
            try:
                export_btn_parent = iframe.locator(
                    '.fr-btn:has(button.x-emb-export)'
                ).first
                cls = await export_btn_parent.get_attribute("class", timeout=1500)
                if cls and "ui-state-enabled" in cls:
                    export_ready = True
            except Exception:
                pass

            if export_ready:
                break
            await asyncio.sleep(poll)
            elapsed += poll
            if elapsed % 30 == 0:
                self.log(f"  仍在等待... ({elapsed}s)", "INFO")

        if export_ready:
            self.log("  报表加载完成", "SUCCESS")
        else:
            self.log("  等待超时，继续尝试导出...", "WARN")
        await asyncio.sleep(1)

    # ================================================================
    # 附件2后处理：删除年累计为0的行 + 指定科目行
    # ================================================================
    def _postprocess_adjustment_excel(self, file_path: str):
        """按 PPT 要求清理网页版调整期国库收支存明细表。"""
        self.log("  开始调整期报表清理...")
        try:
            wb = load_workbook(file_path)
            ws = wb.active
        except Exception as e:
            self.log(f"  读取导出文件失败，跳过清理: {e}", "ERROR")
            return

        try:
            header_row = self._find_header_row(ws)
            if header_row is None:
                self.log("    未识别到表头，跳过调整期清理", "WARN")
                return

            cumulative_col = self._find_cumulative_col(ws, header_row)
            subject_cols = self._find_subject_cols(ws, header_row)
            if cumulative_col is None:
                self.log("    未找到'年累计'列，无法删除年累计为0的行", "WARN")
            if not subject_cols:
                self.log("    未找到科目代码列，将通过整行扫描识别 T040401/T050401", "WARN")

            zero_removed = 0
            subject_removed = 0
            rows_to_delete = []
            for r in range(header_row + 1, ws.max_row + 1):
                remove_for_zero = (
                    cumulative_col is not None
                    and self._is_zero_amount(ws.cell(row=r, column=cumulative_col).value)
                )
                remove_for_subject = self._row_has_excluded_subject(ws, r, subject_cols)
                if remove_for_zero or remove_for_subject:
                    rows_to_delete.append(r)
                    if remove_for_zero:
                        zero_removed += 1
                    if remove_for_subject:
                        subject_removed += 1

            for r in reversed(rows_to_delete):
                ws.delete_rows(r, 1)

            if rows_to_delete:
                wb.save(file_path)
            self.log(
                f"    清理完成：删除 {len(rows_to_delete)} 行"
                f"（年累计=0：{zero_removed}，指定科目：{subject_removed}）",
                "SUCCESS",
            )
        except Exception as e:
            self.log(f"  调整期清理失败: {e}", "ERROR")
            self.log(traceback.format_exc(), "ERROR")
        finally:
            wb.close()

    def _find_header_row(self, ws):
        """在前若干行中寻找包含关键字段的表头行。"""
        max_scan = min(ws.max_row, 20)
        for r in range(1, max_scan + 1):
            headers = [self._norm_header(ws.cell(row=r, column=c).value)
                       for c in range(1, ws.max_column + 1)]
            has_cumulative = any(h == "年累计" for h in headers)
            has_subject = any("科目" in h for h in headers)
            if has_cumulative or has_subject:
                return r
        return None

    def _find_cumulative_col(self, ws, header_row: int):
        for c in range(1, ws.max_column + 1):
            header = self._norm_header(ws.cell(row=header_row, column=c).value)
            if header == "年累计":
                return c
        return None

    def _find_subject_cols(self, ws, header_row: int):
        cols = []
        for c in range(1, ws.max_column + 1):
            header = self._norm_header(ws.cell(row=header_row, column=c).value)
            if "科目" in header and ("代码" in header or "编码" in header):
                cols.append(c)
        return cols

    def _row_has_excluded_subject(self, ws, row_idx: int, subject_cols):
        scan_cols = subject_cols or range(1, ws.max_column + 1)
        for c in scan_cols:
            value = ws.cell(row=row_idx, column=c).value
            if self._norm_subject(value) in EXCLUDED_SUBJECT_CODES:
                return True
        return False

    @staticmethod
    def _norm_header(value):
        return re.sub(r"\s+", "", str(value or "")).strip()

    @staticmethod
    def _norm_subject(value):
        return re.sub(r"\s+", "", str(value or "")).strip().upper()

    @staticmethod
    def _is_zero_amount(value):
        if value is None:
            return False
        if isinstance(value, (int, float)):
            return Decimal(str(value)) == 0
        text = str(value).strip()
        if not text:
            return False
        text = text.replace(",", "").replace("，", "").replace(" ", "")
        if text in {"-", "--"}:
            return False
        if text.startswith("(") and text.endswith(")"):
            text = "-" + text[1:-1]
        try:
            return Decimal(text) == 0
        except (InvalidOperation, ValueError):
            return False

    # ================================================================
    # 导出：导出 → Excel → 原样导出
    # ================================================================
    async def _export_and_save(self, page: Page, dl_folder: str, fname: str) -> str:
        """
        在帆软 iframe 内执行三步导出：
        1. 点击工具栏「导出」按钮
        2. 点击弹出菜单中的「Excel」
        3. 点击子菜单中的「原样导出」
        同时用 expect_download 拦截实际下载。
        """
        self.log("  准备导出...")
        iframe = page.frame_locator(f'iframe[name="{IFRAME_NAME}"]')

        async with page.expect_download(timeout=360000) as dl_info:
            # ----------------------------------------------------------------
            # 步骤 1：点击工具栏「导出」按钮
            # 实际 HTML：<div class="fr-btn ui-state-enabled">
            #               <button class="fr-btn-text x-emb-export">导出</button>
            #            </div>
            # 注意：该 div 没有 widgetname 属性，用内部 button class 定位
            # ----------------------------------------------------------------
            export_btn = iframe.locator('button.fr-btn-text.x-emb-export').first
            await export_btn.wait_for(state="visible", timeout=15000)
            self.log("  点击导出按钮...")
            await export_btn.click()
            await asyncio.sleep(0.8)

            # ----------------------------------------------------------------
            # 步骤 2：点击一级菜单「Excel」
            # 实际 HTML：<div class="fr-ui-core-menu menu">
            #               <div class="menu-item"><div class="menu-text">Excel</div></div>
            #            </div>
            # 菜单在 iframe 内渲染，用 iframe.locator
            # ----------------------------------------------------------------
            excel_item = iframe.locator(
                '.fr-ui-core-menu.menu .menu-item:has(.menu-text:text("Excel"))'
            ).first
            await excel_item.wait_for(state="visible", timeout=5000)
            self.log("  点击 Excel 菜单项...")
            await excel_item.hover()      # hover 触发子菜单展开
            await asyncio.sleep(0.5)

            # ----------------------------------------------------------------
            # 步骤 3：点击二级菜单「原样导出」
            # 实际 HTML（初始 display:none，hover Excel 后显示）：
            #   <div class="fr-ui-core-menu menu">
            #     <div class="menu-item"><div class="menu-text">原样导出</div></div>
            #   </div>
            # ----------------------------------------------------------------
            raw_item = iframe.locator(
                '.fr-ui-core-menu.menu .menu-item:has(.menu-text:text("原样导出"))'
            ).first
            await raw_item.wait_for(state="visible", timeout=5000)
            self.log("  点击原样导出...")
            await raw_item.click()

        download = await dl_info.value
        # 取原始文件扩展名（xlsx / xls），忽略系统生成的文件名主体
        orig = download.suggested_filename or "report.xlsx"
        ext = os.path.splitext(orig)[1] or ".xlsx"
        safe_fname = re.sub(r'[\\/:*?"<>|]', '', fname)
        save_path = os.path.join(dl_folder, f"{safe_fname}{ext}")
        await download.save_as(save_path)
        return save_path


# ============================================================================
# 入口
# ============================================================================
def main():
    app = GDMXApp()
    app.log("调整期国库收支存明细表  批量下载工具  就绪", "SUCCESS")
    app.log("")
    app.log("使用步骤：")
    app.log(f"  1. 输入国库代码（默认 {DEFAULT_TREASURY_CODE}）")
    app.log("  2. 选择保存目录")
    app.log(f"  3. 确认调整期年报年份（默认 {DEFAULT_START_YEAR} ~ {DEFAULT_END_YEAR}）")
    app.log("  4. 点击「捕获登录链接」，在客户端触发系统登录")
    app.log("  5. 程序自动导航到「固定报表→国库收支存明细表」，批量执行全部任务")
    app.log("  6. 如勾选数据处理，导出后自动删除年累计为0的行，以及 T040401/T050401 科目行")
    app.log("")
    app.log(f"  任务量预估：{DEFAULT_END_YEAR - DEFAULT_START_YEAR + 1} 年 × 1库 × 全辖 = {DEFAULT_END_YEAR - DEFAULT_START_YEAR + 1} 次下载", "WARN")
    try:
        app.mainloop()
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()
