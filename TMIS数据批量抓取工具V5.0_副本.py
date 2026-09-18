#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
TMIS数据自由查询批量抓取工具 v5.3
===========================
使用 Playwright 自动启动系统 Chrome 浏览器，携 Token 登录内网系统，
自动填表、查询、导出报表，并可选数据后处理。

功能：
1. 独立登录和常驻浏览器；队列结束后保持待命
2. 自动检测系统 Chrome（channel="chrome"），支持手动指定路径
3. 智能/自定义命名（参数表"文件名称"列 或 自动命名）
4. 收入/支出/退库/库存业务分流（侧边栏动态导航、字段复用映射）
5. 可选数据后处理（openpyxl 保留原始格式）
6. 后备日期输入（默认留空，Excel参数优先）
7. 查询成功精确检测（等待"原样导出"按钮可用）
8. 智能标签页管理（类型切换时自动关闭旧标签页，防止 DOM 冲突）
9. 多参数表、运行中追加、失败重试、暂停和本机任务恢复

依赖：pip install playwright pandas openpyxl psutil
      pip install pyperclip  (可选)
      playwright install chromium
"""

import sys
import os
import re
import threading
import tempfile
import tkinter as tk
from tkinter import messagebox
import pandas as pd
from openpyxl import load_workbook
from datetime import datetime
from decimal import Decimal, InvalidOperation
import time
import asyncio

# Playwright
from playwright.async_api import async_playwright, Page
from tmis_ui import QueueWindow
from tmis_runtime import (
    browser_environment, browser_launch_options, bundled_browser_path,
    normalize_query_date, normalize_task_row, option_matches,
    redact_urls, publish_download, FormInteractionError,
)

# URL 狙击手所需
try:
    import psutil
except ImportError:
    psutil = None

try:
    import pyperclip
except ImportError:
    pyperclip = None


# ============================================================================
# 全局停止标志
# ============================================================================
stop_flag = False          # 停止批量任务
sniper_flag = False        # 停止URL狙击手
SPECIAL_ZERO_SUBJECT_CODES = {"23099", "23106", "2310601", "23202", "2320202"}

# ============================================================================
# Excel 中文表头映射词典
# ============================================================================
COLUMN_MAPPING = {
    "报表库选择": "pRptDbType",
    "调整期标志": "pTrimFlag",
    "报表类型": "pRptType",
    "起始日期": "pStartDate",
    "终止日期": "pEndDate",
    "国库选择": "pTreCode",
    "辖属标志": "pGovernFlag",
    "预算级次": "pBdgLevel",
    "征收机关性质": "pTaxOrgProp",
    "预算科目": "pSbtCode",
    "科目级次": "pStatSbtLevel",
    "金额单位": "pAmtUnit",
    "分国库查询": "pByTre",  # 兼容旧模板
    "预算单位": "pBdgOrgCode",
    "银行大类": "pBnkClass",
    "计划单列市/县/乡": "pSigTreArea",
    "展示范围": "pShowScope",
    "退库原因": "pDwbkReason",
    "国库属性": "pTreAttrib",
    "会计科目": "pBookSbt",
    "会计账户": "pBookAcctName",
    "会计账户名称": "pBookAcctName",
}


# ============================================================================
# 自由查询页面配置
# ============================================================================
REPORT_CONFIGS = {
    "收入": {
        "menu_title": "收入数据自由查询",
        "iframe_name": "fineReportTsasRpt6010",
        "dropdown_fields": [
            "pRptDbType", "pTrimFlag", "pRptType", "pGovernFlag",
            "pTaxOrgProp", "pSigTreArea", "pStatSbtLevel", "pShowScope", "pAmtUnit",
        ],
        "date_fields": ["pStartDate", "pEndDate"],
        "text_fields": ["pTreCode", "pBdgLevel", "pSbtCode"],
        "checkbox_fields": ["分序时", "分地区", "分预算级次", "分预算科目", "分征收机关", "是否展示同比"],
        "template_columns": [
            "文件名称", "是否追加日期", "保留原文件名",
            "报表库选择", "调整期标志", "报表类型", "起始日期", "终止日期",
            "国库选择", "辖属标志", "预算级次", "征收机关性质", "预算科目",
            "计划单列市/县/乡", "科目级次", "展示范围", "金额单位",
            "分序时", "分地区", "分预算级次", "分预算科目", "分征收机关", "是否展示同比",
        ],
        "sample": {
            "文件名称": "收入自由查询", "是否追加日期": "1", "保留原文件名": "1",
            "报表库选择": "1 -- 报表库", "调整期标志": "0 -- 正常期", "报表类型": "",
            "起始日期": "202601", "终止日期": "202601",
            "国库选择": "", "辖属标志": "0 -- 全辖", "预算级次": "",
            "征收机关性质": "0000000000 -- 不分征收机关", "预算科目": "",
            "计划单列市/县/乡": "0 -- 不含", "科目级次": "4 -- 目",
            "展示范围": "1 -- 下级", "金额单位": "2 -- 万元",
            "分序时": "1", "分地区": "0", "分预算级次": "1",
            "分预算科目": "1", "分征收机关": "0", "是否展示同比": "0",
        },
    },
    "支出": {
        "menu_title": "支出数据自由查询",
        "iframe_name": "fineReportTsasRpt6020",
        "dropdown_fields": [
            "pRptDbType", "pTrimFlag", "pRptType", "pGovernFlag",
            "pBdgOrgCode", "pBnkClass", "pSigTreArea", "pStatSbtLevel", "pShowScope", "pAmtUnit",
        ],
        "date_fields": ["pStartDate", "pEndDate"],
        "text_fields": ["pTreCode", "pBdgLevel", "pSbtCode"],
        "checkbox_fields": [
            "分序时", "分地区", "分预算单位", "分银行大类",
            "分预算级次", "分预算科目", "是否展示同比",
        ],
        "template_columns": [
            "文件名称", "是否追加日期", "保留原文件名",
            "报表库选择", "调整期标志", "报表类型", "起始日期", "终止日期",
            "国库选择", "辖属标志", "预算单位", "银行大类", "预算级次", "预算科目",
            "计划单列市/县/乡", "科目级次", "展示范围", "金额单位",
            "分序时", "分地区", "分预算单位", "分银行大类",
            "分预算级次", "分预算科目", "是否展示同比",
        ],
        "sample": {
            "文件名称": "支出自由查询", "是否追加日期": "1", "保留原文件名": "1",
            "报表库选择": "1 -- 报表库", "调整期标志": "0 -- 正常期", "报表类型": "",
            "起始日期": "202601", "终止日期": "202601",
            "国库选择": "", "辖属标志": "0 -- 全辖",
            "预算单位": "ALL -- ALL", "银行大类": "0 -- 全部",
            "预算级次": "", "预算科目": "",
            "计划单列市/县/乡": "0 -- 不含", "科目级次": "4 -- 目",
            "展示范围": "1 -- 下级", "金额单位": "2 -- 万元",
            "分序时": "1", "分地区": "0", "分预算单位": "0", "分银行大类": "0",
            "分预算级次": "1", "分预算科目": "1", "是否展示同比": "0",
        },
    },
    "退库": {
        "menu_title": "退库数据自由查询",
        "iframe_name": "fineReportTsasRpt6030",
        "dropdown_fields": [
            "pRptDbType", "pTrimFlag", "pRptType", "pGovernFlag",
            "pTaxOrgProp", "pSigTreArea", "pStatSbtLevel", "pAmtUnit",
        ],
        "date_fields": ["pStartDate", "pEndDate"],
        "text_fields": ["pTreCode", "pBdgLevel", "pSbtCode", "pDwbkReason"],
        "checkbox_fields": [
            "分序时", "分地区", "分征收机关", "分预算级次",
            "分退库原因", "分预算科目", "是否展示同比",
        ],
        "template_columns": [
            "文件名称", "是否追加日期", "保留原文件名",
            "报表库选择", "调整期标志", "报表类型", "起始日期", "终止日期",
            "国库选择", "辖属标志", "征收机关性质", "预算级次", "预算科目", "退库原因",
            "计划单列市/县/乡", "科目级次", "金额单位",
            "分序时", "分地区", "分征收机关", "分预算级次",
            "分退库原因", "分预算科目", "是否展示同比",
        ],
        "sample": {
            "文件名称": "退库自由查询", "是否追加日期": "1", "保留原文件名": "1",
            "报表库选择": "1 -- 报表库", "调整期标志": "0 -- 正常期", "报表类型": "",
            "起始日期": "202601", "终止日期": "202601",
            "国库选择": "", "辖属标志": "0 -- 全辖",
            "征收机关性质": "0000000000 -- 不分征收机关",
            "预算级次": "", "预算科目": "", "退库原因": "",
            "计划单列市/县/乡": "0 -- 不含", "科目级次": "4 -- 目", "金额单位": "2 -- 万元",
            "分序时": "1", "分地区": "0", "分征收机关": "0", "分预算级次": "1",
            "分退库原因": "1", "分预算科目": "1", "是否展示同比": "0",
        },
    },
    "库存": {
        "menu_title": "库存数据自由查询",
        "iframe_name": "fineReportTsasRpt6040",
        "dropdown_fields": [
            "pRptDbType", "pRptType", "pGovernFlag", "pTreAttrib",
            "pSigTreArea", "pShowScope", "pAmtUnit",
        ],
        "date_fields": ["pStartDate", "pEndDate"],
        "text_fields": ["pTreCode", "pBdgLevel", "pBookSbt", "pBookAcctName"],
        "checkbox_fields": [
            "分序时", "分地区", "分国库属性", "分预算级次",
            "分会计科目", "分会计账户", "是否展示同比",
        ],
        "template_columns": [
            "文件名称", "是否追加日期", "保留原文件名",
            "报表库选择", "报表类型", "起始日期", "终止日期",
            "国库选择", "辖属标志", "国库属性", "预算级次",
            "会计科目", "会计账户名称", "计划单列市/县/乡",
            "展示范围", "金额单位",
            "分序时", "分地区", "分国库属性", "分预算级次",
            "分会计科目", "分会计账户", "是否展示同比",
        ],
        "sample": {
            "文件名称": "库存自由查询", "是否追加日期": "1", "保留原文件名": "1",
            "报表库选择": "1 -- 报表库", "报表类型": "3 -- 月",
            "起始日期": "202601", "终止日期": "202603",
            "国库选择": "", "辖属标志": "0 -- 全辖",
            "国库属性": "0 -- 全部", "预算级次": "0",
            "会计科目": "271,272,273", "会计账户名称": "ALL",
            "计划单列市/县/乡": "0 -- 不含", "展示范围": "1 -- 下级", "金额单位": "0 -- 元",
            "分序时": "1", "分地区": "0", "分国库属性": "0", "分预算级次": "1",
            "分会计科目": "1", "分会计账户": "1", "是否展示同比": "0",
        },
    },
}


# ============================================================================
# 主应用类
# ============================================================================
class ReportEngine:
    """Pure report operations; instantiated independently of the desktop UI."""

    # ========================================================================
    # 智能/自定义命名
    # ========================================================================
    def _build_filename(self, row: pd.Series, sz_type: str, name_counter: dict = None) -> str:
        """根据命名模式生成文件名前缀。"""
        if name_counter is None:
            name_counter = {}

        start_date = str(row.get("pStartDate", "")).strip()
        end_date = str(row.get("pEndDate", "")).strip()
        if not start_date:
            start_date = self.run_options["start_date"]
        if not end_date:
            end_date = self.run_options["end_date"]

        # 模式1：参数表命名
        if self.run_options["naming_mode"] == "param":
            custom_name = str(row.get("文件名称", "")).strip()
            if custom_name:
                append_date = str(row.get("是否追加日期", "1")).strip().lower()
                if append_date in ("0", "否", "false", "no", "n"):
                    return re.sub(r'[\\/:*?"<>|]', '', custom_name)
                if start_date == end_date or not end_date:
                    date_part = start_date if start_date else ""
                else:
                    date_part = f"{start_date}至{end_date}"
                name = f"{custom_name}_{date_part}" if date_part else custom_name
                return re.sub(r'[\\/:*?"<>|]', '', name)
            else:
                self.log("    参数表未填写'文件名称'，回退到自动命名", "WARN")

        # 模式2：自动命名
        if start_date == end_date or not end_date:
            date_part = start_date if start_date else "未知日期"
        else:
            date_part = f"{start_date}至{end_date}"

        bdg_level = str(row.get("pBdgLevel", "")).strip()
        bdg_code = bdg_level.split("--")[0].strip() if "--" in bdg_level else bdg_level
        scale_part = "全口径" if bdg_code in ("", "0") else "地方"

        by_tre = str(row.get("pByTre", "")).strip()
        if "--" in by_tre:
            tre_name = by_tre.split("--")[-1].strip()
        elif "—" in by_tre:
            tre_name = by_tre.split("—")[-1].strip()
        else:
            chinese_chars = re.findall(r'[\u4e00-\u9fff]+', by_tre)
            tre_name = "".join(chinese_chars) if chinese_chars else by_tre
        if not tre_name:
            tre_code = str(row.get("pTreCode", "")).strip()
            tre_name = tre_code if tre_code else "全国"

        name = f"{date_part}_{scale_part}_{sz_type}_{tre_name}"
        safe_name = re.sub(r'[\\/:*?"<>|]', '', name)

        if safe_name in name_counter:
            name_counter[safe_name] += 1
            safe_name = f"{safe_name}_{name_counter[safe_name]}"
        else:
            name_counter[safe_name] = 0

        return safe_name

    def _should_keep_original_name(self, row: pd.Series) -> bool:
        """是否在自定义文件名前继续拼接系统原始导出文件名。"""
        keep = str(row.get("保留原文件名", "1")).strip().lower()
        return keep not in ("0", "否", "false", "no", "n")

    # ========================================================================
    # 需求4：智能侧边栏操作与标签页管理
    # ========================================================================
    async def _close_current_tab(self, page: Page, sz_type: str):
        """关闭当前的查询标签页，防止多标签页 DOM 冲突"""
        tab_name = REPORT_CONFIGS.get(sz_type, {}).get("menu_title", f"{sz_type}数据自由查询")
        self.log(f"  尝试关闭标签页: {tab_name}...", "INFO")
        try:
            # 找到包含对应文本的 tab 节点，然后再点它里面的 .el-icon-close
            close_btn = page.locator(f"span.tags-view-item:has-text('{tab_name}')").locator(".el-icon-close").first
            if await close_btn.is_visible(timeout=3000):
                await close_btn.click()
                self.log(f"  标签页已关闭", "SUCCESS")
                self.current_nav_type = None  # 重置当前导航状态，强制下一次必须点侧边栏
            else:
                self.log(f"  标签页未显示关闭按钮", "WARN")
        except Exception as e:
            self.log(f"  关闭标签页失败: {e}", "WARN")

    async def _navigate_sidebar_smart(self, page: Page, sz_type: str):
        """
        智能导航侧边栏。根据数据类型动态拼接目标菜单。
        如果当前已经导航到同一类型，跳过。

        导航路径：固定报表 -> 数据自由查询 -> {收入/支出/退库/库存}数据自由查询
        """
        config = REPORT_CONFIGS[sz_type]
        menu_title = config["menu_title"]

        if self.current_nav_type and self.current_nav_type != sz_type:
            self.log(f"  检测到数据类型将切换（{self.current_nav_type} -> {sz_type}），正在关闭旧标签页...", "WARN")
            await self._close_current_tab(page, self.current_nav_type)

        if self.current_nav_type == sz_type:
            self.log(f"  当前已在「{menu_title}」页面，跳过导航", "INFO")
            return

        self.log(f"  导航到「{menu_title}」...")

        # 完整菜单路径
        menu_items = [
            "固定报表",
            "数据自由查询",
            menu_title,
        ]

        for item_text in menu_items:
            self.log(f"    点击菜单: {item_text}")

            # 智能容错：先检查子菜单是否已可见
            target = page.get_by_text(item_text, exact=True).first
            try:
                is_vis = await target.is_visible()
            except Exception:
                is_vis = False

            if item_text in ("固定报表", "数据自由查询"):
                # 对于父级菜单，如果目标子菜单已经可见，可跳过点击
                # 检查下一级菜单是否已可见
                next_idx = menu_items.index(item_text) + 1
                if next_idx < len(menu_items):
                    next_item = page.get_by_text(menu_items[next_idx], exact=True).first
                    try:
                        next_vis = await next_item.is_visible()
                        if next_vis:
                            self.log(f"    子菜单已展开，跳过点击: {item_text}")
                            continue
                    except Exception:
                        pass

            # 等待并点击
            try:
                await target.wait_for(state="visible", timeout=10000)
                await target.click()
                await asyncio.sleep(1)
            except Exception:
                # 二次补偿点击
                self.log(f"    首次点击未响应，执行二次补偿点击: {item_text}", "WARN")
                await asyncio.sleep(0.5)
                try:
                    await target.click(force=True)
                    await asyncio.sleep(1)
                except Exception as e2:
                    self.log(f"    菜单点击失败: {item_text} - {e2}", "ERROR")
                    raise

        self.current_nav_type = sz_type
        self.log(f"  导航完成: {menu_title}", "SUCCESS")

    # ========================================================================
    # 需求4：表单填充（支持收入/支出分流）
    # ========================================================================
    async def _fill_form(self, page: Page, row: pd.Series, sz_type: str):
        """
        填充表单，根据数据类型分流处理字段。

        Args:
            page: Playwright Page 对象
            row: Excel 数据行
            sz_type: "收入"、"支出"、"退库" 或 "库存"
        """
        self.log("  开始填充表单...")
        config = REPORT_CONFIGS[sz_type]

        # ---- 下拉选择字段 ----
        for field_id in config["dropdown_fields"]:
            val = str(row.get(field_id, "")).strip()
            if not val:
                continue
            await self._fill_dropdown(page, field_id, val)

        # ---- 日期字段（Excel 优先，UI 后备） ----
        ui_dates = {
            "pStartDate": self.run_options["start_date"],
            "pEndDate": self.run_options["end_date"],
        }
        for field_id in config["date_fields"]:
            val = str(row.get(field_id, "")).strip()
            if not val:
                val = ui_dates.get(field_id, "")
            if not val:
                continue
            await self._fill_date(page, field_id, normalize_query_date(val, row.get("pRptType", "")))

        # ---- 文本输入字段 ----
        for field_id in config["text_fields"]:
            val = str(row.get(field_id, "")).strip()
            if not val:
                continue
            await self._fill_text_input(page, field_id, val)

        # ---- 复选框（根据数据类型分流） ----
        await self._fill_checkboxes(page, row, sz_type)

        self.log("  表单填充完成", "SUCCESS")

    async def _fill_checkboxes(self, page: Page, row: pd.Series, sz_type: str):
        """
        处理复选框。
        根据当前自由查询页面配置，直接读取对应中文列名的 0/1 值。
        """
        for cb_name in REPORT_CONFIGS[sz_type]["checkbox_fields"]:
            val = str(row.get(cb_name, "")).strip()
            if val:
                if val not in ("0", "1"):
                    raise ValueError(f"复选框 {cb_name} 只能填 0 或 1，收到 {val}")
                await self._set_checkbox(page, cb_name, int(val))


    # ========================================================================
    # 表单字段操作方法
    # ========================================================================
    async def _fill_date(self, page: Page, field_id: str, value: str):
        """填充日期字段（el-date-editor）"""
        self.log(f"    日期 {field_id} = {value}")

        # 精确定位：先找到包含该 label 的 form-item 容器，再在容器内找日期输入框
        form_item = page.locator(f'.el-form-item:visible:has(label[for="{field_id}"])').first
        inp = form_item.locator('.el-date-editor .el-input__inner').first
        await inp.wait_for(state="visible", timeout=10000)

        await inp.click()
        await asyncio.sleep(0.3)
        await inp.press("Meta+a" if sys.platform == "darwin" else "Control+a")
        await asyncio.sleep(0.1)
        await inp.press("Delete")
        await asyncio.sleep(0.1)
        await inp.type(value, delay=50)
        await asyncio.sleep(0.2)
        await inp.press("Enter")
        await inp.press("Tab")
        await asyncio.sleep(0.5)
        if (await inp.input_value()).strip() != value:
            raise ValueError(f"日期 {field_id} 未被页面接受，请检查报表类型和日期格式")

    async def _fill_text_input(self, page: Page, field_id: str, value: str):
        """填充文本输入框，填入后 Enter + Tab 触发 Vue 绑定"""
        self.log(f"    文本 {field_id} = {value}")

        # 精确定位：先找到包含该 label 的 form-item 容器，再在容器内找文本输入框
        form_item = page.locator(f'.el-form-item:visible:has(label[for="{field_id}"])').first
        inp = form_item.locator('.el-input .el-input__inner').first
        await inp.wait_for(state="visible", timeout=10000)

        await inp.click()
        await asyncio.sleep(0.2)
        await inp.press("Meta+a" if sys.platform == "darwin" else "Control+a")
        await asyncio.sleep(0.1)
        await inp.press("Delete")
        await asyncio.sleep(0.1)
        await inp.type(value, delay=30)
        await asyncio.sleep(0.2)
        await inp.press("Enter")
        await asyncio.sleep(0.3)
        await inp.press("Tab")
        await asyncio.sleep(0.3)

    async def _fill_dropdown(self, page: Page, field_id: str, value: str):
        """选择 Element-UI 下拉菜单（精确定位到 form-item 容器内的 el-select）"""
        self.log(f"    下拉 {field_id} = {value}")

        # 精确定位：先找到包含该 label 的 form-item 容器，再在容器内找下拉输入框
        form_item = page.locator(f'.el-form-item:visible:has(label[for="{field_id}"])').first
        dropdown_input = form_item.locator('.el-select .el-input__inner').first
        await dropdown_input.wait_for(state="visible", timeout=10000)
        # Avoid opening the popup when the selected value is already correct.
        if option_matches(value, await dropdown_input.input_value()):
            return
        try:
            await page.keyboard.press("Escape")
            await dropdown_input.click(timeout=10000)
            # A fixed 500 ms sleep used to confuse a delayed/hidden popup with
            # an invalid parameter. Wait for a single, populated live popup.
            popup = page.locator('.el-select-dropdown:visible')
            deadline = time.monotonic() + 10
            while time.monotonic() < deadline:
                if await popup.count() == 1 and await popup.locator('.el-select-dropdown__item:visible').count():
                    break
                await asyncio.sleep(0.1)
            else:
                raise FormInteractionError(f'下拉字段 {field_id} 未正常展开或渲染；请还原浏览器或使用无头模式，参数未判为无效')
            visible_items = popup.locator('.el-select-dropdown__item:visible')
            labels = await visible_items.all_text_contents()
            for i, label in enumerate(labels):
                if option_matches(value, label):
                    await visible_items.nth(i).click(timeout=10000)
                    deadline = time.monotonic() + 3
                    while time.monotonic() < deadline:
                        if option_matches(value, await dropdown_input.input_value()):
                            return
                        await asyncio.sleep(0.1)
                    raise FormInteractionError(f'下拉字段 {field_id} 点击后未选中目标值；已停止本条查询')
            await page.keyboard.press("Escape")
            raise ValueError(f"下拉字段 {field_id} 中不存在选项 {value}，本条不执行查询")
        except ValueError:
            raise
        except FormInteractionError:
            raise
        except Exception as error:
            raise FormInteractionError(f'下拉字段 {field_id} 操作未完成；请检查窗口状态，参数未判为无效') from error

    async def _set_checkbox(self, page: Page, label_text: str, target: int):
        """设置复选框状态（通过 is-checked 类判断）"""
        self.log(f"    复选框 '{label_text}' -> {'勾选' if target == 1 else '取消'}")

        cb_locator = page.locator("label.el-checkbox:visible").filter(
            has_text=re.compile(r"^\s*" + re.escape(label_text) + r"\s*$")
        ).first

        try:
            await cb_locator.wait_for(state="visible", timeout=5000)
        except Exception as error:
            raise ValueError(f"页面缺少已指定的复选框: {label_text}") from error

        class_attr = await cb_locator.get_attribute("class") or ""
        is_checked = "is-checked" in class_attr

        if (target == 1 and not is_checked) or (target == 0 and is_checked):
            await cb_locator.click()
            await asyncio.sleep(0.3)

    # ========================================================================
    # 查询等待：检测"原样导出"按钮可用状态
    # ========================================================================
    async def _click_query_and_wait(self, page: Page, sz_type: str = "收入", timeout_seconds=360):
        """先等本次 form POST 导致 iframe 导航，再等新报表的导出按钮。"""
        self.log("  点击查询按钮...")
        iframe_name = REPORT_CONFIGS[sz_type]["iframe_name"]
        query_btn = page.locator("button.el-button:visible:has-text('查询')").first
        await query_btn.wait_for(state="visible", timeout=5000)
        deadline = time.monotonic() + timeout_seconds
        # 提供的页面使用 form[target=fineReportTsasRpt6040] 提交新报表。
        # 只看 ui-state-enabled 会误把前一条任务仍可用的按钮当作新结果。
        async with page.expect_event(
            "framenavigated", predicate=lambda frame: frame.name == iframe_name,
            timeout=timeout_seconds * 1000,
        ):
            await query_btn.click()

        self.log("  本次报表已开始加载，等待原样导出可用...")
        iframe = page.frame_locator(f'iframe[name="{iframe_name}"]')
        export_indicator = iframe.locator('.fr-btn[widgetname="ExcelO"]')
        last_log = time.monotonic()
        while time.monotonic() < deadline:
            if stop_flag:
                raise InterruptedError("用户已停止任务")
            try:
                class_attr = await export_indicator.get_attribute(
                    "class", timeout=max(1, min(2000, int((deadline - time.monotonic()) * 1000))),
                )
                if class_attr and "ui-state-enabled" in class_attr:
                    self.log("  数据查询完成，导出按钮已可用", "SUCCESS")
                    return
            except Exception:
                pass
            await asyncio.sleep(min(0.5, max(0, deadline - time.monotonic())))
            if time.monotonic() - last_log >= 30:
                self.log("  仍在等待本次报表加载...", "INFO")
                last_log = time.monotonic()
        raise TimeoutError("本次查询超时，本条不导出，避免保存前一条报表")

    # ========================================================================
    # 需求2：强力导出（force=True 穿透透明遮挡层）
    # ========================================================================
    async def _export_and_save(
        self,
        page: Page,
        dl_folder: str,
        dynamic_name: str,
        sz_type: str = "收入",
        keep_original_name: bool = True,
        on_saved=None,
    ) -> str:
        """
        切入 iframe，使用 force=True 点击"原样导出"，保存文件。
        返回保存的绝对路径。
        """
        self.log("  准备导出数据...")

        # 切入帆软报表 iframe（收入6010，支出6020，退库6030，库存6040）
        iframe_name = REPORT_CONFIGS[sz_type]["iframe_name"]
        iframe = page.frame_locator(f'iframe[name="{iframe_name}"]')

        # 使用 get_by_text 精确定位"原样导出"按钮
        export_btn = iframe.get_by_text("原样导出", exact=True)
        await export_btn.wait_for(state="visible", timeout=30000)

        self.log("  已找到导出按钮，force=True 强力点击下载...", "INFO")

        # expect_download 拦截下载 + force=True 穿透透明遮挡层
        async with page.expect_download(timeout=360000) as download_info:
            await export_btn.click(force=True)

        download = await download_info.value
        original_name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', '_', download.suggested_filename or "report.xlsx")

        # 默认保持历史行为：{动态名}_{原文件名}；核对专用模板可关闭，精确输出指定文件名。
        ext = os.path.splitext(original_name)[1] or ".xlsx"
        if keep_original_name:
            new_filename = f"{dynamic_name}_{original_name}"
        else:
            new_filename = f"{dynamic_name}{ext}"
        save_path = os.path.join(dl_folder, new_filename)

        fd, temp_path = tempfile.mkstemp(prefix=".tmis-", suffix=".part", dir=dl_folder)
        os.close(fd)
        try:
            await download.save_as(temp_path)
            if await download.failure() or os.path.getsize(temp_path) == 0:
                raise IOError("下载失败或文件为空")
            # Both paths are on the same filesystem. Atomic create-if-absent:
            # never overwrite an earlier report, even when a retry races a file.
            try:
                publish_download(temp_path, save_path)
                if on_saved is not None:
                    on_saved(save_path)
            except FileExistsError as error:
                raise FileExistsError("输出文件已存在，未覆盖；请先核对或移走已有文件再重试：" + new_filename) from error
        finally:
            if os.path.exists(temp_path):
                os.unlink(temp_path)
            # save_as creates a separate copy. Release Playwright's internal
            # temporary download now, not at the end of a multi-hour session.
            try:
                await asyncio.wait_for(download.delete(), timeout=10)
            except Exception:
                self.log('浏览器下载临时副本暂未释放，将在浏览器回收时清理；已保存报表不受影响', 'WARN')
        self.log(f"  文件已保存: {new_filename}", "SUCCESS")

        return save_path

    # ========================================================================
    # 数据后处理（openpyxl 保留格式）
    # ========================================================================
    def _process_excel_data(self, file_path: str, row: pd.Series, sz_type: str):
        """按附件2要求清理自由查询导出表，保留原始格式。"""
        self.log("  开始数据后处理（按PPT要求清理）...")
        try:
            wb = load_workbook(file_path)
            ws = wb.active
        except Exception as e:
            self.log(f"  读取导出文件失败: {e}", "ERROR")
            raise RuntimeError('后处理无法读取下载文件；原文件保留') from e

        if ws.max_row is None or ws.max_row < 2:
            self.log("  导出文件为空，跳过后处理", "WARN")
            wb.close()
            return

        changed = False
        if sz_type == "库存":
            stock_city_col = self._find_header_col(ws, "所属市国库代码")
            if stock_city_col is None:
                self.log("    未找到'所属市国库代码'列，库存专用清理跳过", "INFO")
            else:
                ws.delete_cols(stock_city_col, 1)
                changed = True
                self.log(f"    已删除库存报表'所属市国库代码'列（原第{stock_city_col}列）", "SUCCESS")
        elif sz_type in ("收入", "支出", "退库"):
            removed = self._remove_zero_special_subject_rows(ws)
            changed = removed > 0
            if removed:
                self.log(f"    已删除 {removed} 行发生额和累计额均为0的特殊科目", "SUCCESS")
            else:
                self.log("    未发现需要删除的零值特殊科目行", "INFO")
        else:
            self.log(f"    {sz_type} 暂无PPT指定后处理规则，跳过", "INFO")

        try:
            if changed:
                # Save beside the owned download, then atomically replace it.
                # A failed save must leave the original report intact.
                fd, processed = tempfile.mkstemp(prefix='.tmis-postprocess-', suffix='.xlsx', dir=os.path.dirname(file_path))
                os.close(fd)
                try:
                    wb.save(processed)
                    os.replace(processed, file_path)
                finally:
                    if os.path.exists(processed):
                        os.unlink(processed)
            self.log("  数据后处理完成", "SUCCESS")
        except Exception as e:
            self.log(f"  保存后处理文件失败: {e}", "ERROR")
            raise RuntimeError('后处理保存失败；原文件保留，可重试后处理步骤') from e
        finally:
            wb.close()

    def _remove_zero_special_subject_rows(self, ws):
        header_row = self._find_free_query_header_row(ws)
        if header_row is None:
            self.log("    未识别到表头，零值特殊科目清理跳过", "WARN")
            return 0

        subject_cols = self._find_subject_code_cols(ws, header_row)
        current_col = self._find_amount_col(
            ws, header_row,
            exact_names={"本期执行数", "本期发生额", "发生额"},
            include_tokens=("本期",),
            exclude_tokens=("同期", "同比", "累计"),
        )
        cumulative_col = self._find_amount_col(
            ws, header_row,
            exact_names={"年累计", "累计额"},
            include_tokens=("累计",),
            exclude_tokens=("同期", "同比"),
        )

        if current_col is None or cumulative_col is None:
            self.log("    未找到发生额或累计额列，零值特殊科目清理跳过", "WARN")
            return 0

        rows_to_delete = []
        for r in range(header_row + 1, ws.max_row + 1):
            if not self._row_has_special_subject(ws, r, subject_cols):
                continue
            if (
                self._is_zero_amount(ws.cell(row=r, column=current_col).value)
                and self._is_zero_amount(ws.cell(row=r, column=cumulative_col).value)
            ):
                rows_to_delete.append(r)

        for r in reversed(rows_to_delete):
            ws.delete_rows(r, 1)
        return len(rows_to_delete)

    def _find_free_query_header_row(self, ws):
        max_scan = min(ws.max_row, 20)
        for r in range(1, max_scan + 1):
            headers = [self._norm_header(ws.cell(row=r, column=c).value)
                       for c in range(1, ws.max_column + 1)]
            has_subject = any("科目" in h for h in headers)
            has_current = any(h in {"本期执行数", "本期发生额", "发生额"} for h in headers)
            has_cumulative = any(h == "年累计" or h == "累计额" for h in headers)
            if has_subject and (has_current or has_cumulative):
                return r
        return None

    def _find_subject_code_cols(self, ws, header_row: int):
        cols = []
        for c in range(1, ws.max_column + 1):
            header = self._norm_header(ws.cell(row=header_row, column=c).value)
            if "科目" in header and ("代码" in header or "编码" in header):
                cols.append(c)
        return cols

    def _row_has_special_subject(self, ws, row_idx: int, subject_cols):
        scan_cols = subject_cols or range(1, ws.max_column + 1)
        for c in scan_cols:
            value = self._norm_subject(ws.cell(row=row_idx, column=c).value)
            if value in SPECIAL_ZERO_SUBJECT_CODES:
                return True
        return False

    def _find_amount_col(self, ws, header_row: int, exact_names, include_tokens, exclude_tokens):
        for c in range(1, ws.max_column + 1):
            header = self._norm_header(ws.cell(row=header_row, column=c).value)
            if header in exact_names:
                return c

        for c in range(1, ws.max_column + 1):
            header = self._norm_header(ws.cell(row=header_row, column=c).value)
            if include_tokens and not all(token in header for token in include_tokens):
                continue
            if any(token in header for token in exclude_tokens):
                continue
            return c
        return None

    @staticmethod
    def _norm_header(value):
        return re.sub(r"\s+", "", str(value or "")).strip()

    @staticmethod
    def _norm_subject(value):
        return re.sub(r"\s+", "", str(value or "")).strip().upper()

    @staticmethod
    def _is_zero_amount(value):
        if value is None:
            return False
        if isinstance(value, (int, float)):
            return Decimal(str(value)) == 0
        text = str(value).strip()
        if not text:
            return False
        text = text.replace(",", "").replace("，", "").replace(" ", "")
        if text in {"-", "--"}:
            return False
        if text.startswith("(") and text.endswith(")"):
            text = "-" + text[1:-1]
        try:
            return Decimal(text) == 0
        except (InvalidOperation, ValueError):
            return False

    @staticmethod
    def _find_header_col(ws, header_name: str):
        for c in range(1, ws.max_column + 1):
            header = str(ws.cell(row=1, column=c).value or "").strip()
            if header == header_name:
                return c
        return None


# ============================================================================
class TMISAutoApp(QueueWindow, ReportEngine, tk.Tk):
    """Desktop UI with a persistent service; source configurations stay compatible."""
    report_configs = REPORT_CONFIGS
    column_mapping = COLUMN_MAPPING
    engine_class = ReportEngine

    def _start_sniper(self):
        """启动 URL 狙击手守护线程"""
        global sniper_flag

        if sys.platform.startswith("linux"):
            self._paste_login_url()
            return

        if psutil is None:
            messagebox.showerror("缺少依赖", "请先安装 psutil: pip install psutil")
            return

        if self._state.get('current') or self._state['mode'] in ('running', 'logging_in', 'pausing'):
            messagebox.showwarning("提示", "请先暂停并等待当前条完成，再重新登录", parent=self)
            return

        sniper_flag = False
        self.sniper_btn.config(state=tk.DISABLED, text="  正在监控...  ", bg="#5c6bc0")
        self.log("URL 狙击手已启动 (10ms 极速轮询)", "WARN")
        self.log("请在客户端触发系统登录，本工具将自动拦截 Token 链接", "WARN")

        self.sniper_thread = threading.Thread(target=self._sniper_worker, daemon=True)
        self.sniper_thread.start()

    def _sniper_worker(self):
        """
        URL 狙击手核心逻辑：
        以 10ms 极限间隔高频轮询新浏览器进程，从 cmdline 中提取 http 链接，
        立即 kill 进程（在请求发出前截断），将链接交给独立登录入口。
        此兼容入口只用于非 Linux；不会自动导入或启动批量任务。
        """
        global sniper_flag

        # 记录启动前已有的浏览器进程 PID，避免误杀
        browser_names = {"chrome.exe", "msedge.exe", "firefox.exe", "iexplore.exe",
                         "chrome", "msedge", "firefox", "chromium", "chromium-browser"}
        known_pids = set()
        for proc in psutil.process_iter(["pid", "name"]):
            try:
                if proc.info["name"] and proc.info["name"].lower() in browser_names:
                    known_pids.add(proc.info["pid"])
            except (psutil.NoSuchProcess, psutil.AccessDenied):
                pass

        self.log(f"已标记 {len(known_pids)} 个已有浏览器进程，开始监控...", "INFO")

        captured_url = None
        max_wait = 180  # 最多等待3分钟
        start_time = time.time()

        while not sniper_flag and not self._closing and (time.time() - start_time) < max_wait:
            try:
                for proc in psutil.process_iter(["pid", "name", "cmdline"]):
                    try:
                        pname = (proc.info["name"] or "").lower()
                        if pname not in browser_names:
                            continue
                        if proc.info["pid"] in known_pids:
                            continue

                        # 发现新浏览器进程！扫描命令行参数
                        cmdline = proc.info["cmdline"] or []
                        for arg in cmdline:
                            if isinstance(arg, str) and arg.startswith("http"):
                                captured_url = arg
                                # 立即击杀！在请求发出前掐断进程
                                try:
                                    proc.kill()
                                    self.log(f"已击杀浏览器进程 PID={proc.info['pid']}", "SUCCESS")
                                except Exception as kill_err:
                                    self.log(f"击杀进程失败: {kill_err}", "ERROR")
                                break

                        if captured_url:
                            break

                        known_pids.add(proc.info["pid"])

                    except (psutil.NoSuchProcess, psutil.AccessDenied):
                        continue

                if captured_url:
                    break

            except Exception as e:
                self.log(f"狙击手监控异常: {e}", "ERROR")

            time.sleep(0.01)  # 10ms 极限轮询间隔

        # ---- 处理结果 ----
        if captured_url:
            self.log("Token 链接捕获成功!", "SUCCESS")

            # 同时复制到剪贴板（备用）
            try:
                if pyperclip:
                    pyperclip.copy(captured_url)
                    self.log("链接已同步复制到剪贴板（备用）", "INFO")
            except Exception:
                pass

            self.log("即将打开浏览器登录；参数表和执行队列可另行选择", "SUCCESS")

            # 主线程恢复按钮状态，随后打开独立登录会话。
            self.post_ui(lambda: self.sniper_btn.config(
                state=tk.NORMAL, text="  捕获登录链接  ", bg=self.CLR_BLUE
            ))
            self.post_ui(lambda: self._paste_login_url(captured_url))

        else:
            self.log("狙击手监控超时（180秒），未捕获到链接", "WARN")
            self.post_ui(lambda: self.sniper_btn.config(
                state=tk.NORMAL, text="  捕获登录链接  ", bg=self.CLR_BLUE
            ))




# Excel参数模板生成
# ============================================================================
def generate_template(output_path: str = "TMIS数据自由查询参数模板_v5.1.xlsx"):
    """生成数据自由查询参数模板（分收入/支出/退库/库存 Sheet，全中文表头）"""
    with pd.ExcelWriter(output_path) as writer:
        for report_type, config in REPORT_CONFIGS.items():
            columns = config["template_columns"]
            sample = config["sample"]
            df = pd.DataFrame([sample], columns=columns)
            df.to_excel(writer, sheet_name=f"{report_type}参数", index=False)

    print(f"参数模板已生成: {output_path}")


# ============================================================================
# 入口
# ============================================================================
def main():
    if len(sys.argv) > 1 and sys.argv[1] == "--template":
        out = sys.argv[2] if len(sys.argv) > 2 else "TMIS数据自由查询参数模板_v5.1.xlsx"
        generate_template(out)
        return

    if len(sys.argv) > 1 and sys.argv[1] == "--self-check":
        run_self_check()
        return

    try:
        app = TMISAutoApp()
    except Exception as error:
        print("启动失败：" + redact_urls(error), file=sys.stderr)
        raise SystemExit(1) from error

    app.log("TMIS 数据自由查询批量抓取工具 v5.3 就绪", "SUCCESS")
    app.log("")
    app.log("使用步骤（登录和选表顺序不限）")
    app.log("  1  粘贴链接并登录，可先打开浏览器，再准备参数表")
    app.log("  2  多选参数表，选择保存根目录，点击加入队列")
    app.log("  3  点击开始 / 继续；运行中仍可加入新参数表")
    app.log("  4  本轮结束后可在失败列表重试；浏览器保持开启")
    app.log("  5  暂停等待当前条完成；只有退出程序才关闭浏览器")
    app.log("")
    app.log("参数表日期优先；Linux 可自动使用随包附带的 Chromium。", "INFO")

    try:
        app.mainloop()
    except KeyboardInterrupt:
        pass
    finally:
        app._closing = True
        app.service.shutdown()
        app.service.thread.join(timeout=20)
        if not app.service.thread.is_alive():
            app.store.close()


def run_self_check():
    """在 CI 虚拟桌面验证打包后的 Tk 和 Playwright，不访问 TMIS。"""
    import json
    import platform
    with tempfile.TemporaryDirectory(prefix="tmis-self-check-") as state_dir:
        app = TMISAutoApp(state_dir=state_dir)
        try:
            app.withdraw()
            app.update()
            if sys.platform.startswith("linux"):
                assert not app.sniper_btn.winfo_manager()
            assert callable(app._paste_login_url)
            assert app.task_tree.winfo_exists() and app.failure_tree.winfo_exists()
            assert app.service.thread.is_alive()
        finally:
            app.service.shutdown()
            app.service.thread.join(timeout=20)
            assert not app.service.thread.is_alive()
            app.store.close()
            app.destroy()

    async def check_browser():
        async with async_playwright() as pw:
            browser_path = bundled_browser_path()
            if not browser_path.is_file():
                from pathlib import Path
                browser_path = Path(pw.chromium.executable_path)
            launch_options = {'executable_path': str(browser_path)} if browser_path.is_file() else browser_launch_options()
            browser = await pw.chromium.launch(
                **launch_options, headless=True, env=browser_environment(),
            )
            page = await browser.new_page()
            await page.set_content("<title>TMIS packaging check</title><p>库存自由查询</p>")
            assert await page.title() == "TMIS packaging check"
            await browser.close()

    asyncio.run(check_browser())
    print(json.dumps({"status": "ok", "version": "5.3", "machine": platform.machine(), "libc": platform.libc_ver(), "checks": ["Tk queue GUI", "manual URL entry", "SQLite queue", "persistent service", "Playwright driver", "Chromium launch"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
